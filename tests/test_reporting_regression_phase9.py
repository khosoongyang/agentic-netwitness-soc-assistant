"""Phase 9 of the Reporting redesign: dedicated regression pass, closing
the remaining gaps from the original request's 20-point checklist that
weren't already exercised as a direct side effect of Phases 3-8:

  #10/11 Export Word/PDF exports ONLY the selected report
  #14    Report A can never leak into case B (and: run 1 can never leak
         into run 2 of the SAME incident)
  #18    Missing/incomplete data is handled safely (a few more edge cases)
  #19    No prototype/hardcoded values in the Reporting frontend
  #20    No "Regenerate" action anywhere in the Reporting table

Everything else on the original checklist (four report types, per-report
status/metadata, view/edit/save/review persistence, last-updated/updated-by,
export-uses-saved-version, final-report-from-reviewed-sections, approval
behaviour, existing tests passing) is already covered by
tests/test_report_versions_phase3.py, test_final_report_assembly_phase5.py,
test_candidate_materialisation_phase6.py, test_approval_status_phase7.py,
and the Phase 8 additions to test_reporting_stage.py.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from docx import Document

from workflow import state_store as wss
from agents.reporting import report_editing as re


@pytest.fixture(autouse=True)
def _isolated_db(tmp_path, monkeypatch):
    monkeypatch.setattr(wss, "DB_FILE", tmp_path / "phase9.db")
    wss.db_init()


REPO_ROOT = Path(__file__).resolve().parent.parent


def _blocks(text):
    return [{"type": "paragraph", "text": text}]


def _current_attempt(report_set_id="set-1"):
    return {"report_set_id": report_set_id, "reports": []}


# =============================================================================
# #14 — cross-incident and cross-run isolation
# =============================================================================

def test_report_content_is_isolated_between_incidents():
    re.save_report_edit("INC-A", "INC-A@run1", "executive_summary",
                        _blocks("Incident A content — SECRET-A"), "Analyst A", source_report_set_id="set-a")
    re.save_report_edit("INC-B", "INC-B@run1", "executive_summary",
                        _blocks("Incident B content — SECRET-B"), "Analyst B", source_report_set_id="set-b")

    row_a = re.report_row_state("INC-A", "INC-A@run1", "executive_summary",
                                current_attempt=_current_attempt(), reporting_status="Pending",
                                reporting_updated_at=None)
    row_b = re.report_row_state("INC-B", "INC-B@run1", "executive_summary",
                                current_attempt=_current_attempt(), reporting_status="Pending",
                                reporting_updated_at=None)

    text_a, text_b = json.dumps(row_a["blocks"]), json.dumps(row_b["blocks"])
    assert "SECRET-A" in text_a and "SECRET-B" not in text_a
    assert "SECRET-B" in text_b and "SECRET-A" not in text_b

    versions_a = wss.list_report_versions("INC-A", "INC-A@run1", "executive_summary")
    versions_b = wss.list_report_versions("INC-B", "INC-B@run1", "executive_summary")
    assert all("SECRET-B" not in v["content_json"] for v in versions_a)
    assert all("SECRET-A" not in v["content_json"] for v in versions_b)


def test_report_content_is_isolated_between_runs_of_the_same_incident():
    """A rerun (new run_id, same incident_id) must never see the previous
    run's saved edits, reviews, or content — matching how
    reporting_attempt_dir()/report_versions both key on (incident_id, run_id)."""
    v1 = re.save_report_edit("INC-X", "INC-X@run1", "executive_summary",
                             _blocks("run1 content"), "Analyst A", source_report_set_id="set-1")
    re.mark_reviewed("INC-X", "INC-X@run1", "executive_summary", "Analyst A",
                     expected_report_version_id=v1["id"])
    re.save_report_edit("INC-X", "INC-X@run2", "executive_summary",
                        _blocks("run2 content"), "Analyst B", source_report_set_id="set-2")

    row_run1 = re.report_row_state("INC-X", "INC-X@run1", "executive_summary",
                                   current_attempt=_current_attempt("set-1"), reporting_status="Pending",
                                   reporting_updated_at=None)
    row_run2 = re.report_row_state("INC-X", "INC-X@run2", "executive_summary",
                                   current_attempt=_current_attempt("set-2"), reporting_status="Pending",
                                   reporting_updated_at=None)

    assert row_run1["status"] == "Reviewed"
    assert row_run2["status"] == "Edited"  # run2's own edit, never inherits run1's review
    assert row_run1["blocks"] != row_run2["blocks"]
    assert len(wss.list_report_versions("INC-X", "INC-X@run1", "executive_summary")) == 1
    assert len(wss.list_report_versions("INC-X", "INC-X@run2", "executive_summary")) == 1


def test_report_sets_lineage_is_isolated_between_incidents():
    wss.create_report_set("INC-A", "INC-A@run1", report_set_id="set-a-1", based_on_report_set_id=None,
                          manifest_path="/a", manifest_sha256="aaa", report_version_ids={},
                          status="approved", created_by="X")
    wss.create_report_set("INC-B", "INC-B@run1", report_set_id="set-b-1", based_on_report_set_id=None,
                          manifest_path="/b", manifest_sha256="bbb", report_version_ids={},
                          status="approved", created_by="X")

    assert wss.get_latest_report_set("INC-A", "INC-A@run1")["report_set_id"] == "set-a-1"
    assert wss.get_latest_report_set("INC-B", "INC-B@run1")["report_set_id"] == "set-b-1"
    assert [s["report_set_id"] for s in wss.list_report_sets("INC-A", "INC-A@run1")] == ["set-a-1"]
    assert [s["report_set_id"] for s in wss.list_report_sets("INC-B", "INC-B@run1")] == ["set-b-1"]


# =============================================================================
# #10/11 — export contains ONLY the selected report's content
# =============================================================================

def test_export_report_contains_only_the_selected_reports_content(tmp_path, monkeypatch):
    from workflow import engine as sw
    trusted_root = tmp_path / "trusted_outputs"
    monkeypatch.setattr(sw, "_TRUSTED_OUTPUT_ROOT", trusted_root)

    row_state_a = {"status": "Edited", "blocks": _blocks("Executive Summary UNIQUE TEXT ALPHA"),
                   "title": "Executive Summary", "has_edits": True}
    row_state_b = {"status": "Edited", "blocks": _blocks("Technical Findings UNIQUE TEXT BETA"),
                   "title": "Technical Findings", "has_edits": True}

    data_a, _ = re.export_report("INC-EXPORT", "INC-EXPORT@run1", "executive_summary", "docx",
                                 row_state=row_state_a, reporting_stage_attempt=1, analyst="Analyst A")
    data_b, _ = re.export_report("INC-EXPORT", "INC-EXPORT@run1", "technical_findings", "docx",
                                 row_state=row_state_b, reporting_stage_attempt=1, analyst="Analyst A")

    text_a = "\n".join(p.text for p in Document(io.BytesIO(data_a)).paragraphs)
    text_b = "\n".join(p.text for p in Document(io.BytesIO(data_b)).paragraphs)

    assert "ALPHA" in text_a and "BETA" not in text_a
    assert "BETA" in text_b and "ALPHA" not in text_b


# =============================================================================
# #18 — missing/incomplete data handled safely
# =============================================================================

def test_report_row_state_handles_empty_current_attempt():
    row = re.report_row_state("INC-EMPTY", "INC-EMPTY@run1", "executive_summary",
                              current_attempt={}, reporting_status="Pending", reporting_updated_at=None)
    assert row["status"] == "Not generated"
    assert row["blocks"] == []
    assert row["exists"] is False


def test_list_report_versions_empty_for_report_with_no_versions():
    assert wss.list_report_versions("INC-EMPTY", "INC-EMPTY@run1", "executive_summary") == []


def test_list_report_sets_empty_for_run_with_no_sets():
    assert wss.list_report_sets("INC-EMPTY", "INC-EMPTY@run1") == []
    assert wss.get_latest_report_set("INC-EMPTY", "INC-EMPTY@run1") is None


def test_export_report_raises_cleanly_when_nothing_to_export():
    row_state = {"status": "Not generated", "blocks": [], "title": "Executive Summary", "has_edits": False}
    with pytest.raises(FileNotFoundError):
        re.export_report("INC-EMPTY", "INC-EMPTY@run1", "executive_summary", "docx",
                         row_state=row_state, reporting_stage_attempt=1, analyst="Analyst A")


# =============================================================================
# #19/#20 — the Reporting frontend has no hardcoded prototype values and no
# "Regenerate" action. Static source checks so a future edit can't silently
# reintroduce either without a test failing.
# =============================================================================

_FORBIDDEN_PROTOTYPE_VALUES = (
    "INC-53018", "Incident-008", "KELLYWANG", "JORDAND",
    "26 Jul 2026", "Soong Yang Kho",
)


def test_reports_frontend_has_no_hardcoded_prototype_values():
    source = (REPO_ROOT / "frontend" / "js" / "pages" / "reports.js").read_text(encoding="utf-8")
    for token in _FORBIDDEN_PROTOTYPE_VALUES:
        assert token not in source, f"hardcoded prototype value found in reports.js: {token!r}"


def test_report_editor_component_has_no_hardcoded_prototype_values():
    source = (REPO_ROOT / "frontend" / "js" / "components" / "blockEditor.js").read_text(encoding="utf-8")
    for token in _FORBIDDEN_PROTOTYPE_VALUES:
        assert token not in source, f"hardcoded prototype value found in blockEditor.js: {token!r}"


def test_reports_frontend_never_shows_a_regenerate_action():
    source = (REPO_ROOT / "frontend" / "js" / "pages" / "reports.js").read_text(encoding="utf-8")
    assert "Regenerate" not in source
    assert "regenerate" not in source.lower()


def test_reports_css_never_references_a_regenerate_action():
    source = (REPO_ROOT / "frontend" / "css" / "reports.css").read_text(encoding="utf-8")
    assert "regenerate" not in source.lower()
